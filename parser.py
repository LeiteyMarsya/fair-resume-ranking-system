from __future__ import annotations

from io import BytesIO
from pathlib import Path
import hashlib

import fitz
from docx import Document


# ============================================================
# VERSION
# ============================================================

PARSER_VERSION = "parser-0.1"


# ============================================================
# FILE HASHING
# ============================================================

def calculate_sha256(
    file_bytes: bytes
) -> str:
    """
    Generate a SHA-256 hash for the uploaded file.

    This allows the system to detect the exact same
    file being uploaded multiple times.
    """

    return hashlib.sha256(
        file_bytes
    ).hexdigest()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(
    file_bytes: bytes
) -> str:
    """
    Extract selectable text from a PDF.

    Important:
    This does NOT perform OCR yet.

    If a PDF is a scanned image, very little/no text
    may be returned.
    """

    text_parts: list[str] = []

    document = fitz.open(
        stream=file_bytes,
        filetype="pdf"
    )

    try:

        for page in document:

            page_text = page.get_text(
                "text"
            )

            if page_text:
                text_parts.append(
                    page_text
                )

    finally:

        document.close()

    return "\n".join(
        text_parts
    ).strip()


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(
    file_bytes: bytes
) -> str:
    """
    Extract text from DOCX paragraphs
    and tables.
    """

    document = Document(
        BytesIO(file_bytes)
    )

    parts: list[str] = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                value = cell.text.strip()

                if value:
                    row_values.append(value)

            if row_values:

                parts.append(
                    " | ".join(row_values)
                )

    return "\n".join(
        parts
    ).strip()


# ============================================================
# TXT TEXT EXTRACTION
# ============================================================

def extract_text_from_txt(
    file_bytes: bytes
) -> str:
    """
    Extract text from a TXT file.
    """

    return file_bytes.decode(
        "utf-8",
        errors="replace"
    ).strip()


# ============================================================
# MAIN EXTRACTION FUNCTION
# ============================================================

def extract_text(
    file_bytes: bytes,
    filename: str
) -> tuple[str, str]:
    """
    Automatically choose the correct parser
    based on file extension.

    Returns:

        extracted_text
        extraction_status
    """

    extension = Path(
        filename
    ).suffix.lower()

    try:

        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        if extension == ".pdf":

            text = extract_text_from_pdf(
                file_bytes
            )

        # ----------------------------------------------------
        # DOCX
        # ----------------------------------------------------

        elif extension == ".docx":

            text = extract_text_from_docx(
                file_bytes
            )

        # ----------------------------------------------------
        # TXT
        # ----------------------------------------------------

        elif extension == ".txt":

            text = extract_text_from_txt(
                file_bytes
            )

        # ----------------------------------------------------
        # Unsupported
        # ----------------------------------------------------

        else:

            return (
                "",
                "UNSUPPORTED_FILE_TYPE"
            )

        # ----------------------------------------------------
        # No text
        # ----------------------------------------------------

        if not text:

            return (
                "",
                "NO_TEXT_FOUND"
            )

        # ----------------------------------------------------
        # Very little text
        # ----------------------------------------------------

        if len(text.strip()) < 100:

            return (
                text,
                "LOW_TEXT_WARNING"
            )

        # ----------------------------------------------------
        # Success
        # ----------------------------------------------------

        return (
            text,
            "SUCCESS"
        )

    except Exception as error:

        return (
            "",
            f"ERROR: {error}"
        )